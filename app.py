"""AI Resume ATS Checker - Streamlit + Gemini Flash.

Upload a resume (PDF, DOCX or TXT), optionally paste a job description, and get
an ATS score with concrete suggestions for improvement.
"""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-3.5-flash"
MAX_RESUME_CHARS = 20000
MAX_JD_CHARS = 8000
MIN_RESUME_CHARS = 150

# Category -> weight (sums to 100). The overall score is computed in code from
# the category scores so it is consistent and explainable.
CATEGORY_WEIGHTS = {
    "keywords_relevance": 25,
    "impact_achievements": 25,
    "formatting_structure": 20,
    "skills_presentation": 15,
    "clarity_grammar": 15,
}
CATEGORY_LABELS = {
    "keywords_relevance": "Keywords & Relevance",
    "impact_achievements": "Impact & Achievements",
    "formatting_structure": "Formatting & Structure",
    "skills_presentation": "Skills Presentation",
    "clarity_grammar": "Clarity & Grammar",
}


# --------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------
def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a PDF, DOCX or TXT file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = "\n".join(pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT.")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# --------------------------------------------------------------------------
# Prompt + response handling
# --------------------------------------------------------------------------
def build_prompt(resume_text: str, job_description: str) -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n\"\"\"\n{job_description[:MAX_JD_CHARS]}\n\"\"\"\n"
        if job_description.strip()
        else "No job description was provided. Judge the resume against general "
        "ATS best practices for the role the resume appears to target.\n"
    )
    return f"""You are an expert ATS (Applicant Tracking System) analyst and technical recruiter.
Evaluate the resume below strictly and honestly. Do not inflate scores.

Score each category from 0 to 100:
- keywords_relevance: relevant role keywords/skills, match to the job description if given
- impact_achievements: quantified results, strong action verbs, outcomes over duties
- formatting_structure: standard section headings, ATS-friendly layout, logical order, length
- skills_presentation: clear, categorized, relevant skills section
- clarity_grammar: concise, error-free, consistent tense and style

Return ONLY valid JSON with exactly this shape:
{{
  "candidate_role": "short guess of the role targeted",
  "category_scores": {{
    "keywords_relevance": 0,
    "impact_achievements": 0,
    "formatting_structure": 0,
    "skills_presentation": 0,
    "clarity_grammar": 0
  }},
  "summary": "2-3 sentence overall assessment",
  "strengths": ["..."],
  "improvements": [
    {{"priority": "high|medium|low", "issue": "...", "fix": "..."}}
  ],
  "missing_keywords": ["..."],
  "rewrite_suggestions": [
    {{"original": "a weak line copied from the resume", "improved": "a stronger rewrite"}}
  ]
}}

Rules: give 3-6 strengths, 5-8 improvements ordered by priority, up to 10 missing
keywords, and 2-4 rewrite suggestions. Never invent experience or numbers the
candidate did not state; in rewrites use placeholders like [X%] where a metric is needed.

{jd_block}
RESUME:
\"\"\"
{resume_text[:MAX_RESUME_CHARS]}
\"\"\"
"""


def _clamp(value, low=0, high=100) -> int:
    try:
        return max(low, min(high, int(round(float(value)))))
    except (TypeError, ValueError):
        return 0


def parse_response(raw: str) -> dict:
    """Parse and normalise the model's JSON output."""
    text = (raw or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", text, flags=re.DOTALL)
        if not match:
            raise ValueError("The AI response was not valid JSON.")
        data = json.loads(match.group(0))
    if not isinstance(data, dict):
        raise ValueError("The AI response had an unexpected format.")

    raw_scores = data.get("category_scores") or {}
    scores = {key: _clamp(raw_scores.get(key, 0)) for key in CATEGORY_WEIGHTS}
    overall = round(
        sum(scores[k] * w for k, w in CATEGORY_WEIGHTS.items()) / sum(CATEGORY_WEIGHTS.values())
    )

    def str_list(key):
        items = data.get(key) or []
        return [str(i) for i in items if str(i).strip()] if isinstance(items, list) else []

    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "priority": str(item.get("priority", "medium")).lower(),
                    "issue": str(item.get("issue", "")),
                    "fix": str(item.get("fix", "")),
                }
            )
    order = {"high": 0, "medium": 1, "low": 2}
    improvements.sort(key=lambda i: order.get(i["priority"], 1))

    rewrites = []
    for item in data.get("rewrite_suggestions") or []:
        if isinstance(item, dict) and item.get("original") and item.get("improved"):
            rewrites.append({"original": str(item["original"]), "improved": str(item["improved"])})

    return {
        "overall": overall,
        "category_scores": scores,
        "candidate_role": str(data.get("candidate_role", "")),
        "summary": str(data.get("summary", "")),
        "strengths": str_list("strengths"),
        "improvements": improvements,
        "missing_keywords": str_list("missing_keywords"),
        "rewrite_suggestions": rewrites,
    }


def analyze_resume(resume_text: str, job_description: str, api_key: str, model: str) -> dict:
    """Call Gemini and return the normalised analysis (retries once on bad JSON)."""
    from google import genai
    from google.genai import types

    client = genai.Client(api_key=api_key)
    prompt = build_prompt(resume_text, job_description)
    config = types.GenerateContentConfig(
        response_mime_type="application/json",
        temperature=0.2,
    )
    last_error = None
    for _ in range(2):
        response = client.models.generate_content(model=model, contents=prompt, config=config)
        try:
            return parse_response(response.text)
        except (ValueError, json.JSONDecodeError) as exc:
            last_error = exc
    raise ValueError(f"Could not read the AI response: {last_error}")


# --------------------------------------------------------------------------
# UI
# --------------------------------------------------------------------------
def get_api_key() -> str:
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:
        key = ""
    return key or os.environ.get("GEMINI_API_KEY", "")


def score_color(score: int) -> str:
    if score >= 80:
        return "#16a34a"
    if score >= 60:
        return "#d97706"
    return "#dc2626"


def render_results(result: dict) -> None:
    overall = result["overall"]
    color = score_color(overall)
    verdict = (
        "Strong - likely to pass most ATS filters"
        if overall >= 80
        else "Decent - a few fixes will help noticeably"
        if overall >= 60
        else "Needs work - significant improvements recommended"
    )
    st.markdown(
        f"<div style='text-align:center;padding:1rem'>"
        f"<div style='font-size:4rem;font-weight:700;color:{color}'>{overall}<span style='font-size:1.5rem'>/100</span></div>"
        f"<div style='font-size:1.1rem'>{verdict}</div></div>",
        unsafe_allow_html=True,
    )
    if result["summary"]:
        st.info(result["summary"])

    st.subheader("Score breakdown")
    for key, label in CATEGORY_LABELS.items():
        value = result["category_scores"][key]
        st.write(f"**{label}** ({CATEGORY_WEIGHTS[key]}% weight): {value}/100")
        st.progress(value / 100)

    col1, col2 = st.columns(2)
    with col1:
        st.subheader("Strengths")
        for item in result["strengths"] or ["No strengths returned."]:
            st.markdown(f"- {item}")
    with col2:
        st.subheader("Missing keywords")
        if result["missing_keywords"]:
            st.markdown(" ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.write("None identified.")

    st.subheader("Improvements")
    icons = {"high": "🔴", "medium": "🟠", "low": "🟢"}
    for item in result["improvements"]:
        icon = icons.get(item["priority"], "🟠")
        with st.expander(f"{icon} {item['issue']}", expanded=item["priority"] == "high"):
            st.write(item["fix"])

    if result["rewrite_suggestions"]:
        st.subheader("Suggested rewrites")
        for item in result["rewrite_suggestions"]:
            st.markdown(f"**Before:** {item['original']}")
            st.markdown(f"**After:** {item['improved']}")
            st.divider()


def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="centered")
    st.title("📄 AI Resume ATS Checker")
    st.caption("Upload your resume to get an ATS score and specific ways to improve it.")

    with st.sidebar:
        st.header("Settings")
        api_key = get_api_key()
        if not api_key:
            api_key = st.text_input("Gemini API key", type="password", help="Get one free at aistudio.google.com")
        else:
            st.success("API key loaded")
        model = st.text_input("Gemini model", value=os.environ.get("GEMINI_MODEL", DEFAULT_MODEL))
        st.caption("Your resume is sent to Google's Gemini API for analysis and is not stored by this app.")

    uploaded = st.file_uploader("Resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional)",
        height=150,
        placeholder="Paste the job posting for a tailored keyword match...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please add your Gemini API key in the sidebar.")
            return
        try:
            resume_text = extract_text(uploaded.name, uploaded.getvalue())
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return
        if len(resume_text) < MIN_RESUME_CHARS:
            st.error(
                "Very little text could be extracted. If your resume is a scanned image, "
                "export it as a text-based PDF or DOCX - ATS systems can't read images either."
            )
            return
        try:
            with st.spinner("Analyzing your resume..."):
                result = analyze_resume(resume_text, job_description, api_key, model.strip() or DEFAULT_MODEL)
        except Exception as exc:
            st.error(f"Analysis failed: {exc}")
            return
        st.session_state["result"] = result

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
